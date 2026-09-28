"""Exportación genérica de tablas a XLSX."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_COLOR = "366092"
FOUND_COLOR = "D4EDDA"
NOT_FOUND_COLOR = "F8D7DA"


def _fill(color):
    return PatternFill(start_color=color, end_color=color, fill_type="solid")


def write_table(file_path, sheet_title, headers, rows, widths=None, row_color=None):
    """Guarda `rows` en un XLSX con cabecera formateada.

    `row_color(valores)` puede devolver un color hexadecimal para la fila.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title

    header_fill, header_font = _fill(HEADER_COLOR), Font(bold=True, color="FFFFFF")
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    fills = {}
    for row_num, values in enumerate(rows, 2):
        color = row_color(values) if row_color else None
        fill = fills.setdefault(color, _fill(color)) if color else None
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row_num, column=col, value=str(value))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if fill:
                cell.fill = fill

    for col, width in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    wb.save(file_path)
